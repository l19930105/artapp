# encoding=utf-8
# 导入饼图相关类
from openpyxl.chart import PieChart,Reference
from openpyxl import load_workbook
import os
os.chdir(r"E:\PycharmProjects\SeleniumTest\PythonWork\xiaoqiangSelenium\PortProject\reports")

# 生成饼形图
def pie_chart(wb):
    table = wb.get_sheet_by_name(wb.get_sheet_names()[0])

    # 生成饼图对象
    pie = PieChart()
    #  饼图的标题
    pie.title = "接口测试统计"
    # 获取标签 获得的是范围
    labels = Reference(table, min_col=7, min_row=11, max_col=7, max_row=13)
    # 获取数据，需要注意的是excel中的data上面需要留个空行
    data = Reference(table, min_col=8, min_row=10, max_col=8, max_row=13)

    # 添加数据和标签到图表里
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)

    # 添加图表到sheet里
    table.add_chart(pie, "C21")

    #保存excel
    wb.save("test1.xlsx")

if __name__ == '__main__':
    wb = load_workbook("testcase_report.xlsx")
    pie_chart(wb)
